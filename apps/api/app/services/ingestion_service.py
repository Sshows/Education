from __future__ import annotations

import hashlib
from datetime import datetime
from pathlib import Path

import httpx
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.models import SourceDocument

RAW_DIR = Path("storage/raw")
RAW_DIR.mkdir(parents=True, exist_ok=True)


def _guess_extension(url: str) -> str:
    u = url.lower()
    if ".pdf" in u:
        return ".pdf"
    if ".xlsx" in u or ".xls" in u:
        return ".xlsx"
    return ".html"


def fetch_source(db: Session, source_id: int) -> SourceDocument:
    src = db.scalar(select(SourceDocument).where(SourceDocument.id == source_id))
    if not src:
        raise ValueError("Source not found")

    try:
        r = httpx.get(src.url, timeout=45, follow_redirects=True)
        r.raise_for_status()
        ext = _guess_extension(src.url)
        file_path = RAW_DIR / f"source_{source_id}{ext}"
        file_path.write_bytes(r.content)

        src.file_path = str(file_path)
        src.content_hash = hashlib.sha256(r.content).hexdigest()
        src.fetched_at = datetime.utcnow()
        src.status = "fetched"
        src.error_message = None
        db.commit()
    except Exception as exc:
        src.status = "failed"
        src.error_message = str(exc)
        db.commit()
        raise
    return src


def parse_source(db: Session, source_id: int) -> SourceDocument:
    src = db.scalar(select(SourceDocument).where(SourceDocument.id == source_id))
    if not src:
        raise ValueError("Source not found")
    if not src.file_path:
        raise ValueError("Source file is not fetched")

    path = Path(src.file_path)
    try:
        if path.suffix == ".html":
            text = BeautifulSoup(path.read_text(encoding='utf-8', errors='ignore'), "lxml").get_text(separator=" ", strip=True)
            src.raw_text = text[:200000]
            src.status = "parsed"
        elif path.suffix in {".xlsx", ".xls"}:
            wb = load_workbook(path, read_only=True)
            first = wb[wb.sheetnames[0]]
            rows = []
            for i, row in enumerate(first.iter_rows(values_only=True)):
                if i > 200:
                    break
                rows.append(" | ".join([str(c) for c in row if c is not None]))
            src.raw_text = "\n".join(rows)
            src.status = "parsed"
        elif path.suffix == ".pdf":
            src.status = "failed"
            src.error_message = "PDF parser is not configured yet. Add pypdf/pdfplumber parser in next iteration."
        else:
            src.status = "failed"
            src.error_message = f"Unsupported extension: {path.suffix}"
        db.commit()
    except Exception as exc:
        src.status = "failed"
        src.error_message = str(exc)
        db.commit()
        raise
    return src
